import os
import tempfile
from typing import List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from models import QueryItem

class Exporter:
    """Exports a list of QueryItem entries into an Excel file."""

    # PUBLIC_INTERFACE
    def export(self, items: List[QueryItem]) -> str:
        """
        Export items into an xlsx file and return the absolute file path.

        Columns:
            A: User
            B: Thread URL
            C: Resolver
            D: Status
            E: Created At
            F: Resolved At
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Queries"

        headers = ["User", "Thread URL", "Resolver", "Status", "Created At", "Resolved At"]
        ws.append(headers)

        for it in items:
            ws.append([
                it.user or "",
                it.thread_url or "",
                it.resolver or "",
                it.status or "",
                it.created_at or "",
                it.resolved_at or "",
            ])

        # Auto-size columns (simple heuristic)
        for col in range(1, len(headers) + 1):
            max_length = 0
            col_letter = get_column_letter(col)
            for cell in ws[col_letter]:
                try:
                    val_len = len(str(cell.value)) if cell.value else 0
                    if val_len > max_length:
                        max_length = val_len
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(10, max_length + 2), 80)

        # Save to temp file
        tmp_dir = tempfile.gettempdir()
        out_path = os.path.join(tmp_dir, "slack-queries-export.xlsx")
        wb.save(out_path)
        return out_path
